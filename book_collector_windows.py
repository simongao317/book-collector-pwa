# -*- coding: utf-8 -*-
"""
BookCollector Windows/Python 单文件版

运行:
    python book_collector_windows.py

推荐安装:
    pip install pillow openpyxl opencv-python pyzbar

说明:
    - 数据保存在同目录 book_collector.db
    - CSV 导入导出无需额外依赖
    - XLSX 导入导出需要 openpyxl
    - 摄像头扫码/条码需要 opencv-python + pyzbar
    - 图书信息只从豆瓣图书 book.douban.com/isbn/{ISBN}/ 查询
"""

from __future__ import annotations

import csv
import html
import os
import re
import sqlite3
import sys
import threading
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Callable, Optional

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from PIL import Image, ImageTk
except Exception:  # Pillow 是可选依赖
    Image = None
    ImageTk = None


APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "book_collector.db"


IDENTIFIER_KINDS = ["ISBN", "统一书号", "自定义书号"]
OWNERSHIP_STATUSES = ["已买", "未买"]
OWNED_READING_STATUSES = ["已读", "未读"]
WISHLIST_READING_STATUSES = ["已读", "想读", "待定"]
ENTRY_SOURCES = ["CSV/Excel导入", "扫码导入", "手写导入"]

CSV_COLUMNS = [
    "书号类型",
    "ISBN",
    "统一书号",
    "自定义书号",
    "书名",
    "原名",
    "丛书名",
    "作者",
    "作者国籍",
    "出版时间",
    "出版社",
    "收藏状态",
    "已买阅读标签",
    "未买阅读标签",
    "录入方式",
    "封面路径",
    "添加时间",
    "待定开始时间",
]


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def short_date(value: str) -> str:
    try:
        return datetime.fromisoformat(value).strftime("%y.%m.%d")
    except Exception:
        return datetime.now().strftime("%y.%m.%d")


def normalize_isbn(value: str) -> str:
    return "".join(ch for ch in value.upper() if ch.isdigit() or ch == "X")


def clean_text(value: str) -> str:
    value = html.unescape(value or "")
    value = re.sub(r"<[^>]+>", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def first_match(pattern: str, text: str) -> str:
    match = re.search(pattern, text, flags=re.I | re.S)
    if not match:
        return ""
    return clean_text(match.group(1))


def first_non_empty(*values: str) -> str:
    for value in values:
        if value and value.strip():
            return value.strip()
    return ""


def clean_douban_title(title: str) -> str:
    title = clean_text(title)
    for suffix in [" (豆瓣)", "(豆瓣)", " | 豆瓣", " - 豆瓣"]:
        title = title.replace(suffix, "")
    return title.strip()


def douban_url(isbn: str) -> str:
    return f"https://book.douban.com/isbn/{normalize_isbn(isbn)}/"


def http_get(url: str, timeout: int = 15) -> bytes:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Referer": "https://book.douban.com/",
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def extract_meta(html_text: str, name: str) -> str:
    escaped = re.escape(name)
    patterns = [
        rf"<meta[^>]+(?:property|name)=['\"]{escaped}['\"][^>]+content=['\"]([^'\"]+)['\"][^>]*>",
        rf"<meta[^>]+content=['\"]([^'\"]+)['\"][^>]+(?:property|name)=['\"]{escaped}['\"][^>]*>",
    ]
    for pattern in patterns:
        value = first_match(pattern, html_text)
        if value:
            return value
    return ""


def extract_douban_detail(html_text: str, label: str) -> str:
    label = label.replace(":", "").replace("：", "")
    escaped = re.escape(label)
    patterns = [
        rf"<span[^>]+class=['\"]pl['\"][^>]*>\s*{escaped}\s*:?\s*</span>\s*:?\s*(.*?)<br",
        rf"<span[^>]+class=['\"]pl['\"][^>]*>\s*{escaped}\s*：?\s*</span>\s*：?\s*(.*?)<br",
        rf"{escaped}\s*[:：]\s*([^<\n]+)",
    ]
    for pattern in patterns:
        value = first_match(pattern, html_text)
        if value:
            return value
    return ""


@dataclass
class Book:
    id: Optional[int] = None
    identifier_kind: str = "ISBN"
    isbn: str = ""
    unified_number: str = ""
    custom_number: str = ""
    title: str = ""
    original_title: str = ""
    series_title: str = ""
    authors: str = ""
    author_nationality: str = ""
    publication_date: str = ""
    publisher: str = ""
    ownership_status: str = "已买"
    owned_reading_status: str = "未读"
    wishlist_reading_status: str = "想读"
    entry_source: str = "手写导入"
    cover_path: str = ""
    cover_data: Optional[bytes] = None
    created_at: str = ""
    pending_since: str = ""

    @property
    def primary_identifier(self) -> str:
        if self.identifier_kind == "ISBN":
            return self.isbn.strip()
        if self.identifier_kind == "统一书号":
            return self.unified_number.strip()
        return self.custom_number.strip()

    @property
    def display_reading_status(self) -> str:
        if self.ownership_status == "已买":
            return self.owned_reading_status
        return self.wishlist_reading_status

    @property
    def pending_days_remaining(self) -> Optional[int]:
        if self.ownership_status != "未买" or self.wishlist_reading_status != "待定":
            return None
        start_text = self.pending_since or self.created_at
        try:
            start = datetime.fromisoformat(start_text)
        except Exception:
            start = datetime.now()
        remaining = (start + timedelta(days=30) - datetime.now()).days
        return max(remaining, 0)

    def normalize(self) -> None:
        if self.identifier_kind == "ISBN":
            self.isbn = normalize_isbn(self.isbn)
        elif self.identifier_kind == "统一书号":
            self.unified_number = self.unified_number.strip()
        else:
            self.custom_number = self.custom_number.strip()

        if not self.created_at:
            self.created_at = now_iso()

        if self.ownership_status == "未买" and self.wishlist_reading_status == "待定":
            if not self.pending_since:
                self.pending_since = now_iso()
        else:
            self.pending_since = ""


class LibraryStore:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.create_schema()
        self.cleanup_expired_pending_books()

    def create_schema(self) -> None:
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS books (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                identifier_kind TEXT NOT NULL,
                isbn TEXT,
                unified_number TEXT,
                custom_number TEXT,
                title TEXT,
                original_title TEXT,
                series_title TEXT,
                authors TEXT,
                author_nationality TEXT,
                publication_date TEXT,
                publisher TEXT,
                ownership_status TEXT,
                owned_reading_status TEXT,
                wishlist_reading_status TEXT,
                entry_source TEXT,
                cover_path TEXT,
                cover_data BLOB,
                created_at TEXT,
                pending_since TEXT
            )
            """
        )
        self.conn.commit()

    def all_books(self) -> list[Book]:
        rows = self.conn.execute("SELECT * FROM books ORDER BY id DESC").fetchall()
        return [self.row_to_book(row) for row in rows]

    def add(self, book: Book) -> int:
        book.normalize()
        self.validate(book)
        cur = self.conn.execute(
            """
            INSERT INTO books (
                identifier_kind, isbn, unified_number, custom_number,
                title, original_title, series_title, authors, author_nationality,
                publication_date, publisher, ownership_status, owned_reading_status,
                wishlist_reading_status, entry_source, cover_path, cover_data,
                created_at, pending_since
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            self.book_values(book),
        )
        self.conn.commit()
        return int(cur.lastrowid)

    def update(self, book: Book) -> None:
        book.normalize()
        self.validate(book)
        self.conn.execute(
            """
            UPDATE books SET
                identifier_kind=?, isbn=?, unified_number=?, custom_number=?,
                title=?, original_title=?, series_title=?, authors=?, author_nationality=?,
                publication_date=?, publisher=?, ownership_status=?, owned_reading_status=?,
                wishlist_reading_status=?, entry_source=?, cover_path=?, cover_data=?,
                created_at=?, pending_since=?
            WHERE id=?
            """,
            self.book_values(book) + [book.id],
        )
        self.conn.commit()

    def delete(self, book_id: int) -> None:
        self.conn.execute("DELETE FROM books WHERE id=?", (book_id,))
        self.conn.commit()

    def get(self, book_id: int) -> Book:
        row = self.conn.execute("SELECT * FROM books WHERE id=?", (book_id,)).fetchone()
        if not row:
            raise ValueError("图书不存在")
        return self.row_to_book(row)

    def set_ownership(self, book_id: int, status: str) -> None:
        book = self.get(book_id)
        if book.ownership_status == status:
            return
        if status == "已买" and book.entry_source == "手写导入" and not book.cover_data and not book.cover_path:
            raise ValueError("手写导入已买图书时，需要添加封面照片。")
        book.ownership_status = status
        if status == "已买":
            book.owned_reading_status = "未读"
        else:
            book.wishlist_reading_status = "想读"
        self.update(book)

    def set_reading_status(self, book_id: int, status: str) -> None:
        book = self.get(book_id)
        if book.ownership_status == "已买":
            book.owned_reading_status = status
        else:
            book.wishlist_reading_status = status
        self.update(book)

    def cleanup_expired_pending_books(self) -> None:
        cutoff = datetime.now() - timedelta(days=30)
        for book in self.all_books():
            if book.ownership_status == "未买" and book.wishlist_reading_status == "待定":
                start_text = book.pending_since or book.created_at
                try:
                    start = datetime.fromisoformat(start_text)
                except Exception:
                    start = datetime.now()
                if start < cutoff and book.id is not None:
                    self.delete(book.id)

    def validate(self, book: Book) -> None:
        identifier = book.primary_identifier
        if not identifier:
            raise ValueError("必须填写 ISBN、统一书号或自定义书号之一。")
        rows = self.conn.execute(
            """
            SELECT id FROM books
            WHERE identifier_kind=? AND lower(
                CASE identifier_kind
                    WHEN 'ISBN' THEN isbn
                    WHEN '统一书号' THEN unified_number
                    ELSE custom_number
                END
            )=lower(?)
            """,
            (book.identifier_kind, identifier),
        ).fetchall()
        for row in rows:
            if book.id is None or row["id"] != book.id:
                raise ValueError(f"书号“{identifier}”已经存在，不能重复录入。")

        if book.ownership_status == "已买" and book.entry_source == "手写导入":
            if not book.cover_data and not book.cover_path:
                raise ValueError("手写导入已买图书时，需要添加封面照片。")

    def book_values(self, book: Book) -> list:
        return [
            book.identifier_kind,
            book.isbn,
            book.unified_number,
            book.custom_number,
            book.title,
            book.original_title,
            book.series_title,
            book.authors,
            book.author_nationality,
            book.publication_date,
            book.publisher,
            book.ownership_status,
            book.owned_reading_status,
            book.wishlist_reading_status,
            book.entry_source,
            book.cover_path,
            book.cover_data,
            book.created_at,
            book.pending_since,
        ]

    def row_to_book(self, row: sqlite3.Row) -> Book:
        return Book(
            id=row["id"],
            identifier_kind=row["identifier_kind"] or "ISBN",
            isbn=row["isbn"] or "",
            unified_number=row["unified_number"] or "",
            custom_number=row["custom_number"] or "",
            title=row["title"] or "",
            original_title=row["original_title"] or "",
            series_title=row["series_title"] or "",
            authors=row["authors"] or "",
            author_nationality=row["author_nationality"] or "",
            publication_date=row["publication_date"] or "",
            publisher=row["publisher"] or "",
            ownership_status=row["ownership_status"] or "已买",
            owned_reading_status=row["owned_reading_status"] or "未读",
            wishlist_reading_status=row["wishlist_reading_status"] or "想读",
            entry_source=row["entry_source"] or "手写导入",
            cover_path=row["cover_path"] or "",
            cover_data=row["cover_data"],
            created_at=row["created_at"] or now_iso(),
            pending_since=row["pending_since"] or "",
        )


class DoubanService:
    def lookup(self, isbn: str) -> Book:
        isbn = normalize_isbn(isbn)
        if not isbn:
            raise ValueError("ISBN 不正确，无法查询。")

        url = douban_url(isbn)
        try:
            raw = http_get(url)
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                raise ValueError("没有查询到结果，请手动输入。")
            raise ValueError(f"豆瓣图书返回异常：HTTP {exc.code}")
        except Exception as exc:
            raise ValueError(f"豆瓣图书查询失败：{exc}")

        html_text = raw.decode("utf-8", errors="ignore")
        title = clean_douban_title(
            first_non_empty(
                extract_meta(html_text, "og:title"),
                first_match(r"<title>([^<]+)</title>", html_text),
                first_match(r"<h1[^>]*>(.*?)</h1>", html_text),
            )
        )
        if not title:
            raise ValueError("没有查询到结果，请手动输入。")

        book = Book()
        book.identifier_kind = "ISBN"
        book.isbn = isbn
        book.title = title
        book.original_title = extract_douban_detail(html_text, "原作名")
        book.series_title = extract_douban_detail(html_text, "丛书")
        book.authors = first_non_empty(
            extract_meta(html_text, "book:author"),
            extract_douban_detail(html_text, "作者"),
        )
        book.publisher = extract_douban_detail(html_text, "出版社")
        book.publication_date = extract_douban_detail(html_text, "出版年")
        book.entry_source = "扫码导入"
        book.ownership_status = "已买"
        book.owned_reading_status = "未读"
        book.created_at = now_iso()

        cover_url = extract_meta(html_text, "og:image")
        if cover_url:
            try:
                book.cover_data = http_get(cover_url)
            except Exception:
                pass

        book.normalize()
        return book


class ScannerWindow(tk.Toplevel):
    def __init__(self, master: tk.Tk, on_code: Callable[[str], None]) -> None:
        super().__init__(master)
        self.title("扫码录入")
        self.geometry("420x180")
        self.on_code = on_code
        self.stop_event = threading.Event()

        ttk.Label(self, text="正在打开电脑摄像头，请把 ISBN 条码或二维码对准摄像头。").pack(padx=16, pady=(16, 8))
        ttk.Label(self, text="如果无法打开摄像头，请安装 opencv-python 和 pyzbar。").pack(padx=16, pady=4)
        ttk.Button(self, text="取消", command=self.close).pack(pady=14)

        self.protocol("WM_DELETE_WINDOW", self.close)
        threading.Thread(target=self.scan_loop, daemon=True).start()

    def close(self) -> None:
        self.stop_event.set()
        self.destroy()

    def scan_loop(self) -> None:
        try:
            import cv2
            from pyzbar.pyzbar import decode
        except Exception:
            self.after(0, lambda: messagebox.showerror("缺少依赖", "请先安装：pip install opencv-python pyzbar"))
            self.after(0, self.close)
            return

        cap = cv2.VideoCapture(0)
        if not cap.isOpened():
            self.after(0, lambda: messagebox.showerror("摄像头不可用", "无法打开电脑摄像头。"))
            self.after(0, self.close)
            return

        found = ""
        while not self.stop_event.is_set():
            ok, frame = cap.read()
            if not ok:
                time.sleep(0.05)
                continue
            cv2.imshow("BookCollector - 扫 ISBN 条码/二维码，按 Q 取消", frame)
            for item in decode(frame):
                text = item.data.decode("utf-8", errors="ignore")
                candidate = normalize_isbn(text)
                if len(candidate) in (10, 13):
                    found = candidate
                    break
            if found:
                break
            if cv2.waitKey(1) & 0xFF in (ord("q"), ord("Q")):
                break

        cap.release()
        cv2.destroyAllWindows()
        if found and not self.stop_event.is_set():
            self.after(0, lambda: self.on_code(found))
            self.after(0, self.close)


class BookForm(tk.Toplevel):
    def __init__(self, master: tk.Tk, store: LibraryStore, book: Optional[Book], on_saved: Callable[[], None]) -> None:
        super().__init__(master)
        self.store = store
        self.book = book or Book()
        self.on_saved = on_saved
        self.title("编辑图书" if self.book.id else "录入图书")
        self.geometry("560x720")
        self.resizable(True, True)
        self.vars: dict[str, tk.StringVar] = {}
        self.build()

    def var(self, name: str, value: str = "") -> tk.StringVar:
        variable = tk.StringVar(value=value)
        self.vars[name] = variable
        return variable

    def build(self) -> None:
        frame = ttk.Frame(self, padding=14)
        frame.pack(fill=tk.BOTH, expand=True)

        fields = [
            ("identifier_kind", "书号类型", self.book.identifier_kind, IDENTIFIER_KINDS),
            ("isbn", "ISBN", self.book.isbn, None),
            ("unified_number", "统一书号", self.book.unified_number, None),
            ("custom_number", "自定义书号", self.book.custom_number, None),
            ("title", "书名", self.book.title, None),
            ("original_title", "原名", self.book.original_title, None),
            ("series_title", "丛书名", self.book.series_title, None),
            ("authors", "作者", self.book.authors, None),
            ("author_nationality", "作者国籍", self.book.author_nationality, None),
            ("publication_date", "出版时间", self.book.publication_date, None),
            ("publisher", "出版社", self.book.publisher, None),
            ("ownership_status", "收藏状态", self.book.ownership_status, OWNERSHIP_STATUSES),
            ("owned_reading_status", "已买阅读标签", self.book.owned_reading_status, OWNED_READING_STATUSES),
            ("wishlist_reading_status", "未买阅读标签", self.book.wishlist_reading_status, WISHLIST_READING_STATUSES),
            ("entry_source", "录入方式", self.book.entry_source, ENTRY_SOURCES),
        ]

        for row, (key, label, value, choices) in enumerate(fields):
            ttk.Label(frame, text=label).grid(row=row, column=0, sticky="w", pady=4)
            variable = self.var(key, value)
            if choices:
                widget = ttk.Combobox(frame, textvariable=variable, values=choices, state="readonly")
            else:
                widget = ttk.Entry(frame, textvariable=variable)
            widget.grid(row=row, column=1, sticky="ew", pady=4)

        cover_row = len(fields)
        self.cover_label_var = self.var("cover_path", self.book.cover_path)
        ttk.Label(frame, text="封面路径").grid(row=cover_row, column=0, sticky="w", pady=4)
        ttk.Entry(frame, textvariable=self.cover_label_var).grid(row=cover_row, column=1, sticky="ew", pady=4)
        ttk.Button(frame, text="选择封面", command=self.choose_cover).grid(row=cover_row, column=2, padx=6)

        ttk.Frame(frame).grid(row=cover_row + 1, column=0, columnspan=3, pady=8)
        button_frame = ttk.Frame(frame)
        button_frame.grid(row=cover_row + 2, column=0, columnspan=3, sticky="e")
        ttk.Button(button_frame, text="取消", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(button_frame, text="保存", command=self.save).pack(side=tk.RIGHT, padx=4)

        frame.columnconfigure(1, weight=1)

    def choose_cover(self) -> None:
        path = filedialog.askopenfilename(
            title="选择封面图片",
            filetypes=[("图片文件", "*.jpg *.jpeg *.png *.webp *.bmp"), ("所有文件", "*.*")],
        )
        if path:
            self.cover_label_var.set(path)
            try:
                self.book.cover_data = Path(path).read_bytes()
            except Exception:
                self.book.cover_data = None

    def save(self) -> None:
        for key, variable in self.vars.items():
            setattr(self.book, key, variable.get())
        if self.book.cover_path:
            try:
                self.book.cover_data = Path(self.book.cover_path).read_bytes()
            except Exception:
                pass
        try:
            if self.book.id:
                self.store.update(self.book)
            else:
                self.store.add(self.book)
            self.on_saved()
            self.destroy()
        except Exception as exc:
            messagebox.showerror("无法保存", str(exc))


class BookCollectorApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("BookCollector")
        self.geometry("1100x680")
        self.store = LibraryStore(DB_PATH)
        self.douban = DoubanService()
        self.search_var = tk.StringVar()
        self.filter_var = tk.StringVar(value="全部")
        self.status_var = tk.StringVar(value="全部")
        self.thumbnail_cache: dict[int, object] = {}
        self.build_ui()
        self.refresh()

    def build_ui(self) -> None:
        toolbar = ttk.Frame(self, padding=(10, 8))
        toolbar.pack(fill=tk.X)

        ttk.Label(toolbar, text="筛选").pack(side=tk.LEFT)
        filters = [
            "全部",
            "已买",
            "已买 · 已读",
            "已买 · 未读",
            "未买",
            "未买 · 已读",
            "未买 · 想读",
            "未买 · 待定",
        ]
        filter_box = ttk.Combobox(toolbar, textvariable=self.filter_var, values=filters, state="readonly", width=16)
        filter_box.pack(side=tk.LEFT, padx=6)
        filter_box.bind("<<ComboboxSelected>>", lambda _e: self.refresh())

        ttk.Label(toolbar, textvariable=self.status_var, font=("", 11, "bold")).pack(side=tk.LEFT, padx=18)
        ttk.Button(toolbar, text="全部", command=self.reset_filter).pack(side=tk.LEFT)

        ttk.Entry(toolbar, textvariable=self.search_var, width=28).pack(side=tk.LEFT, padx=(20, 6))
        ttk.Button(toolbar, text="搜索", command=self.refresh).pack(side=tk.LEFT)
        ttk.Button(toolbar, text="清空", command=self.clear_search).pack(side=tk.LEFT, padx=4)

        ttk.Button(toolbar, text="扫码", command=self.scan).pack(side=tk.RIGHT, padx=4)
        ttk.Button(toolbar, text="手写录入", command=self.add_manual).pack(side=tk.RIGHT, padx=4)
        ttk.Button(toolbar, text="导入", command=self.import_file).pack(side=tk.RIGHT, padx=4)
        ttk.Button(toolbar, text="导出", command=self.export_file).pack(side=tk.RIGHT, padx=4)

        columns = ("title", "authors", "publisher", "identifier", "ownership", "reading", "pending", "created")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", selectmode="browse")
        headings = {
            "title": "书名",
            "authors": "作者",
            "publisher": "出版社",
            "identifier": "书号",
            "ownership": "已买/未买",
            "reading": "阅读标签",
            "pending": "待定剩余",
            "created": "添加时间",
        }
        widths = {
            "title": 230,
            "authors": 150,
            "publisher": 190,
            "identifier": 150,
            "ownership": 90,
            "reading": 90,
            "pending": 90,
            "created": 90,
        }
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], anchor=tk.W)
        self.tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 8))
        self.tree.bind("<Double-1>", lambda _e: self.edit_selected())

        actions = ttk.Frame(self, padding=(10, 0, 10, 10))
        actions.pack(fill=tk.X)
        ttk.Button(actions, text="编辑", command=self.edit_selected).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="删除", command=self.delete_selected).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="设为已买", command=lambda: self.set_ownership("已买")).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="设为未买", command=lambda: self.set_ownership("未买")).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="已读", command=lambda: self.set_reading("已读")).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="未读", command=lambda: self.set_reading("未读")).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="想读", command=lambda: self.set_reading("想读")).pack(side=tk.LEFT, padx=4)
        ttk.Button(actions, text="待定", command=lambda: self.set_reading("待定")).pack(side=tk.LEFT, padx=4)

    def refresh(self) -> None:
        self.store.cleanup_expired_pending_books()
        for item in self.tree.get_children():
            self.tree.delete(item)

        filter_name = self.filter_var.get()
        self.status_var.set(filter_name)
        query = self.search_var.get().strip().lower()

        for book in self.store.all_books():
            if not self.matches_filter(book, filter_name):
                continue
            haystack = " ".join(
                [
                    book.title,
                    book.original_title,
                    book.series_title,
                    book.authors,
                    book.publisher,
                    book.primary_identifier,
                ]
            ).lower()
            if query and query not in haystack:
                continue
            pending = ""
            days = book.pending_days_remaining
            if days is not None:
                pending = f"{days}天后删除"
            self.tree.insert(
                "",
                tk.END,
                iid=str(book.id),
                values=(
                    book.title or "未命名图书",
                    book.authors,
                    book.publisher,
                    f"{book.identifier_kind} {book.primary_identifier}",
                    book.ownership_status,
                    book.display_reading_status,
                    pending,
                    short_date(book.created_at),
                ),
            )

    def matches_filter(self, book: Book, filter_name: str) -> bool:
        if filter_name == "全部":
            return True
        if filter_name == "已买":
            return book.ownership_status == "已买"
        if filter_name == "已买 · 已读":
            return book.ownership_status == "已买" and book.owned_reading_status == "已读"
        if filter_name == "已买 · 未读":
            return book.ownership_status == "已买" and book.owned_reading_status == "未读"
        if filter_name == "未买":
            return book.ownership_status == "未买"
        if filter_name == "未买 · 已读":
            return book.ownership_status == "未买" and book.wishlist_reading_status == "已读"
        if filter_name == "未买 · 想读":
            return book.ownership_status == "未买" and book.wishlist_reading_status == "想读"
        if filter_name == "未买 · 待定":
            return book.ownership_status == "未买" and book.wishlist_reading_status == "待定"
        return True

    def reset_filter(self) -> None:
        self.filter_var.set("全部")
        self.refresh()

    def clear_search(self) -> None:
        self.search_var.set("")
        self.refresh()

    def selected_book_id(self) -> Optional[int]:
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo("提示", "请先选择一本书。")
            return None
        return int(selected[0])

    def add_manual(self, draft: Optional[Book] = None) -> None:
        book = draft or Book()
        book.entry_source = "手写导入"
        BookForm(self, self.store, book, self.refresh)

    def edit_selected(self) -> None:
        book_id = self.selected_book_id()
        if book_id is None:
            return
        BookForm(self, self.store, self.store.get(book_id), self.refresh)

    def delete_selected(self) -> None:
        book_id = self.selected_book_id()
        if book_id is None:
            return
        if messagebox.askyesno("确认删除", "确定删除这本书吗？"):
            self.store.delete(book_id)
            self.refresh()

    def set_ownership(self, status: str) -> None:
        book_id = self.selected_book_id()
        if book_id is None:
            return
        try:
            self.store.set_ownership(book_id, status)
            self.refresh()
        except Exception as exc:
            messagebox.showerror("无法修改", str(exc))

    def set_reading(self, status: str) -> None:
        book_id = self.selected_book_id()
        if book_id is None:
            return
        try:
            book = self.store.get(book_id)
            if book.ownership_status == "已买" and status not in OWNED_READING_STATUSES:
                messagebox.showinfo("提示", "已买图书只能设置为“已读”或“未读”。")
                return
            if book.ownership_status == "未买" and status not in WISHLIST_READING_STATUSES:
                messagebox.showinfo("提示", "未买图书只能设置为“已读”“想读”或“待定”。")
                return
            self.store.set_reading_status(book_id, status)
            self.refresh()
        except Exception as exc:
            messagebox.showerror("无法修改", str(exc))

    def scan(self) -> None:
        ScannerWindow(self, self.handle_scanned_code)

    def handle_scanned_code(self, code: str) -> None:
        isbn = normalize_isbn(code)
        if not isbn:
            messagebox.showerror("扫码失败", "没有识别到有效 ISBN。")
            return

        progress = tk.Toplevel(self)
        progress.title("查询中")
        progress.geometry("300x100")
        ttk.Label(progress, text=f"正在用豆瓣查询 ISBN {isbn} ...").pack(expand=True, padx=16, pady=16)

        def worker() -> None:
            try:
                book = self.douban.lookup(isbn)
                self.store.add(book)
                self.after(0, lambda: self.scan_success(progress, book))
            except Exception as exc:
                self.after(0, lambda: self.scan_failed(progress, isbn, exc))

        threading.Thread(target=worker, daemon=True).start()

    def scan_success(self, window: tk.Toplevel, book: Book) -> None:
        window.destroy()
        self.refresh()
        messagebox.showinfo("已保存", f"已添加：{book.title}")

    def scan_failed(self, window: tk.Toplevel, isbn: str, exc: Exception) -> None:
        window.destroy()
        if messagebox.askyesno("没有查询到结果", f"没有查询到结果，请手动输入。\n\n{exc}\n\n是否打开手动录入？"):
            draft = Book(identifier_kind="ISBN", isbn=isbn, entry_source="扫码导入", ownership_status="已买")
            self.add_manual(draft)

    def import_file(self) -> None:
        path = filedialog.askopenfilename(
            title="导入图书数据",
            filetypes=[("CSV/XLSX", "*.csv *.xlsx"), ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not path:
            return
        try:
            suffix = Path(path).suffix.lower()
            rows = self.read_xlsx(path) if suffix == ".xlsx" else self.read_csv(path)
            count = 0
            for row in rows:
                book = self.row_to_book(row)
                book.entry_source = "CSV/Excel导入"
                self.store.add(book)
                count += 1
            self.refresh()
            messagebox.showinfo("导入完成", f"已导入 {count} 本图书。")
        except Exception as exc:
            messagebox.showerror("导入失败", str(exc))

    def export_file(self) -> None:
        path = filedialog.asksaveasfilename(
            title="导出图书数据",
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            suffix = Path(path).suffix.lower()
            rows = [self.book_to_row(book) for book in self.store.all_books()]
            if suffix == ".xlsx":
                self.write_xlsx(path, rows)
            else:
                self.write_csv(path, rows)
            messagebox.showinfo("导出完成", f"已导出到：{path}")
        except Exception as exc:
            messagebox.showerror("导出失败", str(exc))

    def read_csv(self, path: str) -> list[dict[str, str]]:
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def write_csv(self, path: str, rows: list[dict[str, str]]) -> None:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)

    def read_xlsx(self, path: str) -> list[dict[str, str]]:
        try:
            import openpyxl
        except Exception:
            raise RuntimeError("导入 XLSX 需要先安装：pip install openpyxl")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [str(cell.value or "") for cell in ws[1]]
        rows = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            rows.append({headers[i]: str(values[i] or "") for i in range(len(headers))})
        return rows

    def write_xlsx(self, path: str, rows: list[dict[str, str]]) -> None:
        try:
            import openpyxl
        except Exception:
            raise RuntimeError("导出 XLSX 需要先安装：pip install openpyxl")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Books"
        ws.append(CSV_COLUMNS)
        for row in rows:
            ws.append([row.get(col, "") for col in CSV_COLUMNS])
        wb.save(path)

    def row_to_book(self, row: dict[str, str]) -> Book:
        cover_path = row.get("封面路径", "")
        cover_data = None
        if cover_path and Path(cover_path).exists():
            try:
                cover_data = Path(cover_path).read_bytes()
            except Exception:
                cover_data = None
        book = Book(
            identifier_kind=row.get("书号类型") or "ISBN",
            isbn=row.get("ISBN", ""),
            unified_number=row.get("统一书号", ""),
            custom_number=row.get("自定义书号", ""),
            title=row.get("书名", ""),
            original_title=row.get("原名", ""),
            series_title=row.get("丛书名", ""),
            authors=row.get("作者", ""),
            author_nationality=row.get("作者国籍", ""),
            publication_date=row.get("出版时间", ""),
            publisher=row.get("出版社", ""),
            ownership_status=row.get("收藏状态") or "已买",
            owned_reading_status=row.get("已买阅读标签") or "未读",
            wishlist_reading_status=row.get("未买阅读标签") or "想读",
            entry_source=row.get("录入方式") or "CSV/Excel导入",
            cover_path=cover_path,
            cover_data=cover_data,
            created_at=row.get("添加时间") or now_iso(),
            pending_since=row.get("待定开始时间", ""),
        )
        return book

    def book_to_row(self, book: Book) -> dict[str, str]:
        return {
            "书号类型": book.identifier_kind,
            "ISBN": book.isbn,
            "统一书号": book.unified_number,
            "自定义书号": book.custom_number,
            "书名": book.title,
            "原名": book.original_title,
            "丛书名": book.series_title,
            "作者": book.authors,
            "作者国籍": book.author_nationality,
            "出版时间": book.publication_date,
            "出版社": book.publisher,
            "收藏状态": book.ownership_status,
            "已买阅读标签": book.owned_reading_status,
            "未买阅读标签": book.wishlist_reading_status,
            "录入方式": book.entry_source,
            "封面路径": book.cover_path,
            "添加时间": book.created_at,
            "待定开始时间": book.pending_since,
        }


def main() -> None:
    if sys.version_info < (3, 10):
        print("请使用 Python 3.10 或更新版本。")
        return
    app = BookCollectorApp()
    app.mainloop()


if __name__ == "__main__":
    main()
